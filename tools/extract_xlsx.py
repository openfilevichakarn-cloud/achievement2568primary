#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""แปลงไฟล์ .xlsm (ปพ.5/ปพ.6 ป.N) เป็น JSON สำหรับเว็บ: เซลล์ สูตร สไตล์ การผสานเซลล์ ความกว้าง/สูง Data Validation การพิมพ์
ใช้:  python extract_xlsx.py "ไฟล์.xlsm" out.json [--skip unblock,macro]"""
import sys, json, re, zipfile, colorsys, argparse, warnings
import openpyxl
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.utils import get_column_letter, range_boundaries
from lxml import etree
warnings.filterwarnings("ignore")

NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
      "x14": "http://schemas.microsoft.com/office/spreadsheetml/2009/9/main",
      "xm": "http://schemas.microsoft.com/office/excel/2006/main"}

def theme_colors(wb):
    t = etree.fromstring(wb.loaded_theme)
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    cs = t.find(".//a:clrScheme", ns)
    out = []
    for ch in cs:
        c = ch[0]
        out.append(c.get("val") if c.tag.endswith("srgbClr") else c.get("lastClr"))
    # Excel index order: lt1, dk1, lt2, dk2, accent1..6, hlink, folHlink
    order = [out[1], out[0], out[3], out[2]] + out[4:10] + out[10:12]
    return ["#" + x for x in order]

def tint(hexcol, t):
    if not t: return hexcol
    r, g, b = [int(hexcol[i:i+2], 16) / 255 for i in (1, 3, 5)]
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    l = l * (1 + t) if t < 0 else l * (1 - t) + t
    r, g, b = colorsys.hls_to_rgb(h, min(max(l, 0), 1), s)
    return "#%02X%02X%02X" % (round(r * 255), round(g * 255), round(b * 255))

def color(c, th):
    if c is None: return None
    try:
        if c.type == "rgb":
            v = c.rgb
            if isinstance(v, str) and len(v) == 8:
                if v[:2] == "00" and v != "00000000": pass
                return "#" + v[2:]
            return None
        if c.type == "theme":
            return tint(th[c.theme], c.tint)
        if c.type == "indexed":
            v = COLOR_INDEX[c.indexed] if c.indexed < len(COLOR_INDEX) else None
            return "#" + v[2:] if v else None
    except Exception:
        return None
    return None


DNS = {"xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
       "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
WING2 = {"R": "\u2611", "T": "\u2612", "S": "\u2612", "\u00a3": "\u2610", "P": "\u2713", "O": "\u2717"}

def draw_color(el, th):
    """สีใน DrawingML -> #RRGGBB (รองรับ srgb / sysClr / schemeClr แบบหลัก)"""
    if el is None: return None
    c = el[0] if len(el) else None
    if c is None: return None
    tag = etree.QName(c).localname
    if tag == "srgbClr": return "#" + c.get("val")
    if tag == "sysClr": return "#" + (c.get("lastClr") or "000000")
    if tag == "schemeClr":
        v = c.get("val"); m = {"bg1": "#FFFFFF", "lt1": "#FFFFFF", "tx1": "#000000", "dk1": "#000000", "bg2": th[2], "tx2": th[3]}
        return m.get(v) or {"accent1": th[4], "accent2": th[5], "accent3": th[6], "accent4": th[7], "accent5": th[8], "accent6": th[9]}.get(v)
    return None

def read_shapes(z, sheet_path, th):
    """กล่อง/รูปทรงเรียบ ๆ ที่วางทับเซลล์ (เช่น ช่องสี่เหลี่ยมติ๊กถูก) — ข้ามปุ่มมาโคร ลูกศร และกล่องที่มีข้อความ"""
    rp = sheet_path.rsplit("/", 1)[0] + "/_rels/" + sheet_path.rsplit("/", 1)[1] + ".rels"
    if rp not in z.namelist(): return []
    rr = etree.fromstring(z.read(rp)); dpath = None
    for r in rr:
        if r.get("Type", "").endswith("/drawing"):
            t = r.get("Target"); dpath = "xl/" + t.replace("../", "") if t.startswith("..") else t.lstrip("/")
    if not dpath or dpath not in z.namelist(): return []
    d = etree.fromstring(z.read(dpath)); out = []
    for an in d.findall("xdr:twoCellAnchor", DNS):
        sp = an.find("xdr:sp", DNS)
        if sp is None or sp.get("macro"): continue
        if "".join(sp.itertext()).strip(): continue
        spPr = sp.find("xdr:spPr", DNS); geom = spPr.find("a:prstGeom", DNS)
        g = geom.get("prst") if geom is not None else "rect"
        if g not in ("rect", "roundRect", "ellipse"): continue
        def pos(tag):
            e = an.find("xdr:" + tag, DNS)
            return [int(e.find("xdr:" + k, DNS).text) for k in ("col", "colOff", "row", "rowOff")]
        fill = None if spPr.find("a:noFill", DNS) is not None else draw_color(spPr.find("a:solidFill", DNS), th)
        ln = spPr.find("a:ln", DNS); line = None; lw = 1.5
        if ln is not None and ln.find("a:noFill", DNS) is None:
            line = draw_color(ln.find("a:solidFill", DNS), th) or "#000000"
            if ln.get("w"): lw = max(1, round(int(ln.get("w")) / 9525, 1))
        out.append({"from": pos("from"), "to": pos("to"), "g": g, "fill": fill, "line": line, "lw": lw})
    return out

def wingdings_map(z):
    """ข้อความ rich text ที่ใช้ฟอนต์ Wingdings -> ข้อความ Unicode (เช่น 'R ผ่าน £ ไม่ผ่าน' -> '☑ ผ่าน ☐ ไม่ผ่าน')"""
    out = {}
    if "xl/sharedStrings.xml" not in z.namelist(): return out
    sst = etree.fromstring(z.read("xl/sharedStrings.xml")); M = "{%s}" % NS["m"]
    for si in sst.findall(M + "si"):
        runs = si.findall(M + "r")
        if not runs: continue
        plain, conv, hit = "", "", False
        for r in runs:
            t = "".join(x.text or "" for x in r.findall(M + "t"))
            f = r.find(M + "rPr/" + M + "rFont")
            plain += t
            if f is not None and "wingdings" in (f.get("val") or "").lower():
                hit = True; conv += "".join(WING2.get(ch, ch) for ch in t)
            else: conv += t
        if hit: out[plain] = conv
    return out

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsm"); ap.add_argument("out"); ap.add_argument("--skip", default="unblock,macro")
    a = ap.parse_args()
    skip = set(a.skip.split(","))
    wb = openpyxl.load_workbook(a.xlsm)
    wbv = openpyxl.load_workbook(a.xlsm, data_only=True)
    th = theme_colors(wb)
    z = zipfile.ZipFile(a.xlsm)
    wbx = etree.fromstring(z.read("xl/workbook.xml"))
    rels = etree.fromstring(z.read("xl/_rels/workbook.xml.rels"))
    rid2t = {r.get("Id"): r.get("Target") for r in rels}
    sheet_xml = {}
    for s in wbx.find("m:sheets", NS):
        t = rid2t[s.get("{%s}id" % NS["r"])]
        sheet_xml[s.get("name")] = "xl/" + t.lstrip("/").replace("xl/", "") if not t.startswith("/") else t.lstrip("/")
    styles, style_idx = [], {}
    def sid(c):
        f, fl, b, al = c.font, c.fill, c.border, c.alignment
        def side(s):
            return [s.style, color(s.color, th) or "#000000"] if s and s.style else None
        d = {
            "f": [f.name, f.sz, 1 if f.b else 0, 1 if f.i else 0, f.u or 0, color(f.color, th) or "#000000"],
            "bg": color(fl.fgColor, th) if fl and fl.fill_type == "solid" else None,
            "b": [side(b.left), side(b.right), side(b.top), side(b.bottom)],
            "a": [al.horizontal, al.vertical, 1 if al.wrap_text else 0, al.indent or 0, al.textRotation or 0, 1 if al.shrink_to_fit else 0],
            "n": c.number_format,
        }
        k = json.dumps(d, sort_keys=True, ensure_ascii=False)
        if k not in style_idx:
            style_idx[k] = len(styles); styles.append(d)
        return style_idx[k]
    # print areas
    pareas = {}
    for ws in wb.worksheets:
        pa = ws.print_area
        if pa: pareas[ws.title] = pa if isinstance(pa, str) else ",".join(pa)
    sheets = []
    for ws in wb.worksheets:
        if ws.title in skip: continue
        wv = wbv[ws.title]
        cells = {}
        for row in ws.iter_rows():
            for c in row:
                if c.value is None and not c.has_style: continue
                d = {}
                s = sid(c)
                if s: d["s"] = s
                v = c.value
                if isinstance(v, str) and v.startswith("="):
                    d["f"] = v[1:]
                    cv = wv[c.coordinate].value
                    if cv is not None: d["c"] = cv
                elif v is not None:
                    d["v"] = v
                if c.protection.locked is False: d["u"] = 1
                if c.comment: d["cm"] = c.comment.text
                if d: cells[c.coordinate] = d
        # col widths
        cw = {}
        dflt = ws.sheet_format.defaultColWidth or (ws.sheet_format.baseColWidth or 8) + 0.71
        for k, cd in ws.column_dimensions.items():
            lo, hi = cd.min, cd.max
            if lo is None: continue
            for i in range(lo, min(hi, 60) + 1):
                cw[i] = {"w": cd.width, "h": 1 if cd.hidden else 0}
        rh = {}
        for r, rd in ws.row_dimensions.items():
            if rd.height is not None or rd.hidden:
                rh[r] = {"h": rd.height, "x": 1 if rd.hidden else 0}
        dv = []
        for d in (ws.data_validations.dataValidation if ws.data_validations else []):
            dv.append({"t": d.type, "op": d.operator, "f1": d.formula1, "f2": d.formula2, "sq": str(d.sqref),
                       "ab": 1 if d.allow_blank else 0, "et": d.errorTitle, "e": d.error, "ps": d.promptTitle, "p": d.prompt,
                       "sa": 1 if d.showErrorMessage else 0})
        # x14 validations (cross-sheet lists)
        sx = etree.fromstring(z.read(sheet_xml[ws.title]))
        for x in sx.iterfind(".//x14:dataValidation", NS):
            f1 = x.find("x14:formula1/xm:f", NS)
            sq = x.find("xm:sqref", NS)
            dv.append({"t": x.get("type"), "op": x.get("operator"), "f1": f1.text if f1 is not None else None, "f2": None,
                       "sq": sq.text if sq is not None else "", "ab": 1 if x.get("allowBlank") else 0, "et": x.get("errorTitle"),
                       "e": x.get("error"), "sa": 1 if x.get("showErrorMessage") else 0})
        sfp = sx.find("m:sheetFormatPr", NS)
        zero = sfp is not None and sfp.get("zeroHeight") == "1"
        ps = ws.page_setup
        sp = ws.sheet_properties
        sheets.append({
            "name": ws.title, "hidden": ws.sheet_state != "visible",
            "tab": color(sp.tabColor, th) if sp.tabColor else None,
            "dim": [ws.max_row, ws.max_column],
            "dcw": dflt, "drh": ws.sheet_format.defaultRowHeight or 15, "zero": zero,
            "cw": {str(k): v for k, v in cw.items()}, "rh": {str(k): v for k, v in rh.items()},
            "merges": [str(m) for m in ws.merged_cells.ranges],
            "cells": cells, "shapes": ([] if ws.sheet_state != "visible" else read_shapes(z, sheet_xml[ws.title], th)), "dv": dv, "prot": 1 if ws.protection.sheet else 0,
            "grid": 1 if ws.sheet_view.showGridLines is not False else 0,
            "freeze": ws.freeze_panes,
            "print": {"area": pareas.get(ws.title), "orient": ps.orientation or "portrait", "paper": ps.paperSize,
                      "scale": ps.scale, "fitPage": 1 if (sp.pageSetUpPr and sp.pageSetUpPr.fitToPage) else 0,
                      "fw": ps.fitToWidth, "fh": ps.fitToHeight,
                      "m": [ws.page_margins.left, ws.page_margins.right, ws.page_margins.top, ws.page_margins.bottom],
                      "hc": 1 if ws.print_options.horizontalCentered else 0,
                      "breaks": [b.id for b in ws.row_breaks.brk]},
        })
    wmap = wingdings_map(z)
    for sh in sheets:
        for c in sh["cells"].values():
            if isinstance(c.get("v"), str) and c["v"] in wmap: c["v"] = wmap[c["v"]]
    out = {"styles": styles, "sheets": sheets, "source": a.xlsm.split("/")[-1]}
    json.dump(out, open(a.out, "w", encoding="utf-8"), ensure_ascii=False, separators=(",", ":"))
    print(len(sheets), "sheets,", len(styles), "styles,", sum(len(s["cells"]) for s in sheets), "cells")

main()
